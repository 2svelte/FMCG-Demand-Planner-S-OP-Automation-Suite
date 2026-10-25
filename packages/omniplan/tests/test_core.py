import io

import pandas as pd
from openpyxl import load_workbook

from omniplan import forecast, master_data, scenario
from omniplan.excel_export import build_workbook


def test_forecast_shape_and_lift(world):
    pos, master, promo, base, lift = world
    assert len(base) == pos[["SKU_ID", "Retailer"]].drop_duplicates().shape[0] * 12
    assert (base["Baseline_Units"] >= 0).all()
    assert (lift["Promo_Lift"].dropna() > 0).all()
    up = forecast.apply_promos(base, lift, promo, 0.3)["Promo_Uplift_Units"].sum()
    flat = forecast.apply_promos(base, lift, promo, 0.0)["Promo_Uplift_Units"].sum()
    assert up == flat * 1.3 or abs(up - flat * 1.3) < 1e-6


def test_delay_creates_stockouts(world):
    pos, master, promo, base, lift = world
    plan = forecast.apply_promos(base, lift, promo, 0.0)
    _, s0 = scenario.project_inventory(plan, plan, pos, master, 0)
    _, s3 = scenario.project_inventory(plan, plan, pos, master, 3)
    assert (s3["Risk_Status"] != "OK").sum() > (s0["Risk_Status"] != "OK").sum()
    assert s0["Capital_At_Risk"].sum() > 0


def test_master_data_audit_finds_seeded_issues(world):
    pos, master, *_ = world
    a = master_data.audit(master, pos)
    pairs = set(zip(a["SKU_ID"], a["Issue"]))
    assert ("PC-1011", "Discontinued SKU with active demand") in pairs
    assert ("PC-1006", "Missing MOQ") in pairs
    assert ("PC-1009", "Missing lead time") in pairs
    assert ("PC-1003", "Abnormal order multiple") in pairs
    assert ("PC-1002", "Abnormal POS swing (non-promo)") in pairs


def test_excel_has_formulas(world):
    pos, master, promo, base, lift = world
    plan = forecast.apply_promos(base, lift, promo, 0.0)
    wi, summ = scenario.project_inventory(plan, plan, pos, master, 1)
    cons = scenario.consensus_vs_budget(plan, pos, master)
    audit = master_data.audit(master, pos)
    wb = load_workbook(io.BytesIO(build_workbook(plan, wi, summ, cons, audit, pos, master, 0.0, 1, 120)))
    assert {"Assumptions", "Forward Forecast", "Inventory Projection", "S&OP Consensus", "Master Data Issues"} <= set(wb.sheetnames)
    assert str(wb["Forward Forecast"]["Q2"].value).startswith("=SUM(")
    assert str(wb["S&OP Consensus"]["E2"].value) == "=C2*(1+D2)"
